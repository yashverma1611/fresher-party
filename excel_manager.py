"""
Excel Manager module for Event QR Pass & Scanner Management System.
Handles multi-sheet Excel processing for BCA 1st & 3rd Semester.
Provides template download, bulk upload/parsing, and post-event attendance export.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import uuid
import datetime

HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep Navy
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1")
)

CHECKED_IN_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Light Green
PENDING_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Light Amber

def create_sample_template() -> bytes:
    """Generate a clean 2-sheet Excel template for BCA 1st Sem and BCA 3rd Sem."""
    wb = openpyxl.Workbook()
    # Sheet 1: BCA 1st Sem
    ws1 = wb.active
    ws1.title = "BCA 1st Sem"

    headers = ["Roll No", "Student Name", "Semester", "Email Address", "Phone / Mobile"]
    ws1.append(headers)

    sample_1st = [
        ["26BCA101", "Aarav Sharma", "BCA 1st", "aarav.bca26@gmail.com", "+91 9123456701"],
        ["26BCA102", "Ananya Iyer", "BCA 1st", "ananya.iyer@gmail.com", "+91 9123456702"],
        ["26BCA103", "Rohan Mehta", "BCA 1st", "rohan.mehta@gmail.com", "+91 9123456703"],
        ["26BCA104", "Sneha Kapoor", "BCA 1st", "sneha.k@gmail.com", "+91 9123456704"],
        ["26BCA105", "Kabir Sen", "BCA 1st", "kabir.sen@gmail.com", "+91 9123456705"],
    ]
    for row in sample_1st:
        ws1.append(row)

    # Sheet 2: BCA 3rd Sem
    ws2 = wb.create_sheet(title="BCA 3rd Sem")
    ws2.append(headers)

    sample_3rd = [
        ["25BCA301", "Devansh Malhotra", "BCA 3rd", "devansh.m@gmail.com", "+91 9876543101"],
        ["25BCA302", "Pooja Reddy", "BCA 3rd", "pooja.reddy@gmail.com", "+91 9876543102"],
        ["25BCA303", "Siddharth Jain", "BCA 3rd", "sid.jain@gmail.com", "+91 9876543103"],
        ["25BCA304", "Meera Kulkarni", "BCA 3rd", "meera.k@gmail.com", "+91 9876543104"],
        ["25BCA305", "Aditya Deshmukh", "BCA 3rd", "aditya.d@gmail.com", "+91 9876543105"],
    ]
    for row in sample_3rd:
        ws2.append(row)

    for ws in [ws1, ws2]:
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

def parse_uploaded_excel(file_bytes_or_path):
    """
    Parse an uploaded Excel workbook.
    Extracts student records from BCA 1st Sem and BCA 3rd Sem sheets or general sheet.
    Returns list of dicts: [{roll_no, name, semester, email, phone}]
    """
    if isinstance(file_bytes_or_path, bytes):
        wb = openpyxl.load_workbook(filename=BytesIO(file_bytes_or_path), data_only=True)
    else:
        wb = openpyxl.load_workbook(filename=file_bytes_or_path, data_only=True)

    extracted_students = []
    seen_rolls = set()

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        default_sem = "BCA 1st"
        if "3" in sheet_name.lower() or "3rd" in sheet_name.lower():
            default_sem = "BCA 3rd"
        elif "1" in sheet_name.lower() or "1st" in sheet_name.lower():
            default_sem = "BCA 1st"

        # Find header row
        header_row_idx = None
        col_map = {}

        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if not row or all(v is None for v in row):
                continue
            str_row = [str(v or '').strip().lower() for v in row]
            if any('roll' in c for c in str_row) or any('name' in c for c in str_row):
                header_row_idx = r_idx
                for c_idx, cell_val in enumerate(str_row):
                    if 'roll' in cell_val:
                        col_map['roll_no'] = c_idx
                    elif 'name' in cell_val:
                        col_map['name'] = c_idx
                    elif 'sem' in cell_val:
                        col_map['semester'] = c_idx
                    elif 'email' in cell_val:
                        col_map['email'] = c_idx
                    elif 'phone' in cell_val or 'mobile' in cell_val or 'contact' in cell_val:
                        col_map['phone'] = c_idx
                break

        if header_row_idx is None:
            continue

        # Extract data rows
        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if not row or all(v is None for v in row):
                continue

            roll_val = str(row[col_map['roll_no']] or '').strip() if 'roll_no' in col_map and col_map['roll_no'] < len(row) else ''
            name_val = str(row[col_map['name']] or '').strip() if 'name' in col_map and col_map['name'] < len(row) else ''

            if not roll_val or not name_val:
                continue

            sem_val = default_sem
            if 'semester' in col_map and col_map['semester'] < len(row) and row[col_map['semester']]:
                raw_sem = str(row[col_map['semester']]).strip()
                if '3' in raw_sem:
                    sem_val = 'BCA 3rd'
                elif '1' in raw_sem:
                    sem_val = 'BCA 1st'

            email_val = str(row[col_map['email']] or '').strip() if 'email' in col_map and col_map['email'] < len(row) and row[col_map['email']] else f"{roll_val.lower()}@college.edu"
            phone_val = str(row[col_map['phone']] or '').strip() if 'phone' in col_map and col_map['phone'] < len(row) and row[col_map['phone']] else "+91 9999999999"

            if roll_val in seen_rolls:
                continue
            seen_rolls.add(roll_val)

            extracted_students.append({
                'roll_no': roll_val,
                'name': name_val,
                'semester': sem_val,
                'email': email_val,
                'phone': phone_val
            })

    return extracted_students

def export_attendance_excel(conn) -> bytes:
    """Generate professional 4-sheet attendance workbook with audit trail."""
    wb = openpyxl.Workbook()
    cursor = conn.cursor()

    # Sheet 1: Master Attendance
    ws_master = wb.active
    ws_master.title = "Master Attendance"

    # Sheet 2: BCA 1st Sem
    ws_1st = wb.create_sheet(title="BCA 1st Sem")

    # Sheet 3: BCA 3rd Sem
    ws_3rd = wb.create_sheet(title="BCA 3rd Sem")

    # Sheet 4: Audit Logs
    ws_audit = wb.create_sheet(title="Audit Log")

    att_headers = [
        "Roll No", "Student Name", "Semester", "Status",
        "Checked-In Time", "Verified By Admin", "Verification Method",
        "Email", "Phone", "Pass Token"
    ]

    for ws in [ws_master, ws_1st, ws_3rd]:
        ws.append(att_headers)

    cursor.execute("""
    SELECT roll_no, name, semester, status, checked_in_at, scanned_by_admin_name,
           scan_method, email, phone, pass_token
    FROM students
    ORDER BY semester ASC, roll_no ASC
    """)
    all_students = cursor.fetchall()

    for s in all_students:
        status_label = "CHECKED IN" if s["status"] == "checked_in" else "PENDING"
        row_data = [
            s["roll_no"],
            s["name"],
            s["semester"],
            status_label,
            s["checked_in_at"] or "-",
            s["scanned_by_admin_name"] or "-",
            s["scan_method"] or "-",
            s["email"],
            s["phone"],
            s["pass_token"]
        ]
        ws_master.append(row_data)
        if "1st" in s["semester"]:
            ws_1st.append(row_data)
        elif "3rd" in s["semester"]:
            ws_3rd.append(row_data)

    # Style student sheets
    for ws in [ws_master, ws_1st, ws_3rd]:
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        for row_idx in range(2, ws.max_row + 1):
            status_cell = ws.cell(row=row_idx, column=4)
            if status_cell.value == "CHECKED IN":
                status_cell.fill = CHECKED_IN_FILL
                status_cell.font = Font(name="Calibri", size=11, bold=True, color="166534")
            else:
                status_cell.fill = PENDING_FILL
                status_cell.font = Font(name="Calibri", size=11, bold=True, color="92400E")

            for col_idx in range(1, len(att_headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = THIN_BORDER
                if col_idx in [1, 3, 4, 5, 7]:
                    c.alignment = Alignment(horizontal="center", vertical="center")

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    # Populate Audit Log Sheet
    audit_headers = [
        "Log ID", "Timestamp", "Student Name", "Roll No",
        "Semester", "Admin Name", "Method", "Action", "Details"
    ]
    ws_audit.append(audit_headers)
    for cell in ws_audit[1]:
        cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_audit.row_dimensions[1].height = 28

    cursor.execute("""
    SELECT id, timestamp, student_name, roll_no, semester, admin_name, method, action, details
    FROM audit_logs
    ORDER BY id DESC
    """)
    for log in cursor.fetchall():
        ws_audit.append([
            log["id"],
            log["timestamp"],
            log["student_name"],
            log["roll_no"],
            log["semester"],
            log["admin_name"],
            log["method"],
            log["action"],
            log["details"] or ""
        ])

    for row_idx in range(2, ws_audit.max_row + 1):
        for col_idx in range(1, len(audit_headers) + 1):
            c = ws_audit.cell(row=row_idx, column=col_idx)
            c.border = THIN_BORDER

    for col in ws_audit.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_audit.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

if __name__ == "__main__":
    t = create_sample_template()
    print("Template created. Size:", len(t), "bytes")
